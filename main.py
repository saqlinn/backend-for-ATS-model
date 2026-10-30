import os
import shutil
import subprocess
import uuid
from pathlib import Path
from typing import Any

import uvicorn
from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel, ConfigDict

import db
from pipeline import analyze_resume_text

BASE_DIR = Path(__file__).resolve().parent
UPLOAD_DIR = BASE_DIR / "uploads"
MAX_UPLOAD_BYTES = int(os.getenv("MAX_UPLOAD_BYTES", str(10 * 1024 * 1024)))
ALLOWED_EXTENSIONS = {".pdf", ".docx", ".doc"}

app = FastAPI(title="ATS Resume Analyzer API", version="2.0")
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)
db.init_db()


class AnalysisResponse(BaseModel):
    model_config = ConfigDict(extra="allow")
    analysis_id: str | None = None
    status: str
    category: str | None = None
    primary_role: str | None = None
    stated_job_title: str | None = None
    secondary_roles: list[dict[str, Any]] = []
    skills: dict[str, Any] = {}
    strong_skills: list[str] = []
    missing_skills: list[str] = []
    score_breakdown: dict[str, Any] = {}
    recommendations: list[dict[str, Any]] = []
    extraction_warnings: list[str] = []
    limitations: list[str] = []


@app.get("/")
def read_root() -> dict[str, str]:
    return {"message": "ATS Resume Analyzer API is live!"}


@app.get("/health")
def health_check() -> dict[str, str]:
    return {"status": "ok"}


def _extract_pdf(path: Path) -> tuple[str, list[str]]:
    import fitz

    warnings: list[str] = []
    with fitz.open(path) as document:
        if document.is_encrypted:
            raise ValueError("The PDF is encrypted and cannot be analyzed.")
        pages = [page.get_text("text").strip() for page in document]
    text = "\n\n".join(page for page in pages if page)
    if not text:
        try:
            import pytesseract
            from PIL import Image

            with fitz.open(path) as document:
                ocr_pages = []
                for page in document:
                    pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                    image = Image.frombytes("RGB", [pixmap.width, pixmap.height], pixmap.samples)
                    ocr_pages.append(pytesseract.image_to_string(image).strip())
            text = "\n\n".join(page for page in ocr_pages if page)
            if text:
                warnings.append("Text was recovered using OCR; verify extracted details.")
        except (ImportError, OSError, RuntimeError) as error:
            warnings.append(f"PDF has no extractable text and OCR is unavailable: {error}")
    return text, warnings


def _extract_docx(path: Path) -> tuple[str, list[str]]:
    from docx import Document

    document = Document(path)
    parts = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells if cell.text.strip()))
    return "\n".join(parts), []


def _extract_legacy_doc(path: Path) -> tuple[str, list[str]]:
    for command in (["antiword", str(path)], ["soffice", "--headless", "--convert-to", "txt:Text", "--outdir", str(path.parent), str(path)]):
        executable = shutil.which(command[0])
        if executable:
            completed = subprocess.run([executable, *command[1:]], capture_output=True, text=True, timeout=30, check=False)
            if completed.returncode == 0:
                output_path = path.with_suffix(".txt")
                if output_path.exists():
                    return output_path.read_text(encoding="utf-8", errors="replace"), ["Legacy .doc content was converted before analysis."]
                if completed.stdout.strip():
                    return completed.stdout, ["Legacy .doc content was converted before analysis."]
    raise ValueError("Legacy .doc files require antiword or LibreOffice (soffice) to be installed.")


def extract_text_from_file(path: str | Path) -> tuple[str, list[str]]:
    source = Path(path)
    suffix = source.suffix.casefold()
    if suffix == ".pdf":
        return _extract_pdf(source)
    if suffix == ".docx":
        return _extract_docx(source)
    if suffix == ".doc":
        return _extract_legacy_doc(source)
    raise ValueError("Unsupported file type. Upload a PDF, DOCX, or DOC resume.")


async def _save_upload(file: UploadFile, destination: Path) -> int:
    size = 0
    with destination.open("wb") as buffer:
        while chunk := await file.read(1024 * 1024):
            size += len(chunk)
            if size > MAX_UPLOAD_BYTES:
                raise HTTPException(status_code=413, detail=f"File exceeds the {MAX_UPLOAD_BYTES} byte upload limit.")
            buffer.write(chunk)
    return size


@app.post("/v1/analyzeResume", response_model=AnalysisResponse)
async def analyze_resume(user_id: str | None = Form(None), file: UploadFile = File(...)) -> dict[str, Any]:
    original_filename = Path(file.filename or "resume").name
    suffix = Path(original_filename).suffix.casefold()
    if suffix not in ALLOWED_EXTENSIONS:
        raise HTTPException(status_code=415, detail="Unsupported file type. Upload a PDF, DOCX, or DOC resume.")

    UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
    tmp_path = UPLOAD_DIR / f"{uuid.uuid4()}{suffix}"
    try:
        await _save_upload(file, tmp_path)
        try:
            resume_text, extraction_warnings = extract_text_from_file(tmp_path)
        except (ValueError, ImportError, OSError, RuntimeError) as error:
            raise HTTPException(status_code=400, detail=f"Failed to extract resume text: {error}") from error
        if len(resume_text.strip()) < 50:
            raise HTTPException(status_code=422, detail="The document contains too little extractable text to analyze.")
        analysis = analyze_resume_text(resume_text, user_id=user_id, filename=original_filename)
        analysis["extraction_warnings"] = extraction_warnings + analysis.get("extraction_warnings", [])
        db.save_analysis(user_id or "anonymous", original_filename, resume_text, analysis)
        return analysis
    finally:
        tmp_path.unlink(missing_ok=True)


if __name__ == "__main__":
    db.init_db()
    uvicorn.run(app, host="0.0.0.0", port=8000)
